# core/zonal_exporter.py
# Экспорт зональных данных в Excel — компактная сводка на 1 страницу (книжная)
from datetime import datetime
from typing import List, Optional
from .zonal_models import ZonalCollection, ReportItemType, Criminalist
from .zonal_data import get_report_data, get_summary, is_item_filled, get_criminalist_fill
from .constants import INITIAL_DEPARTMENTS


class ZonalExcelExporter:
    """Экспорт данных зональных криминалистов в Excel — одна страница"""

    def export(self, collection: ZonalCollection, filepath: str) -> None:
        """Создать и сохранить .xlsx файл — компактная сводка на 1 лист"""
        print(f"[ZONAL_EXCEL] Начинаю экспорт, криминалистов: {len(collection.criminalists)}")

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
        except ImportError:
            raise ImportError("openpyxl не установлена. pip install openpyxl")

        wb = Workbook()
        ws = wb.active
        ws.title = "Сводка"

        # ── Стили ─────────────────────────────────────────────────
        header_fill = PatternFill("solid", fgColor="1E293B")
        header_font = Font(bold=True, color="FFFFFF", size=10)
        subheader_fill = PatternFill("solid", fgColor="334155")
        subheader_font = Font(bold=True, color="FFFFFF", size=9)
        name_font = Font(bold=True, size=10, color="1E293B")
        data_font = Font(size=10, color="334155")
        bold_data = Font(size=10, bold=True, color="1E293B")
        green_fill = PatternFill("solid", fgColor="DCFCE7")
        green_font = Font(size=10, bold=True, color="15803D")
        yellow_fill = PatternFill("solid", fgColor="FEF9C3")
        yellow_font = Font(size=10, bold=True, color="92400E")
        red_fill = PatternFill("solid", fgColor="FEE2E2")
        red_font = Font(size=10, color="991B1B")
        gray_fill = PatternFill("solid", fgColor="F1F5F9")
        summary_fill = PatternFill("solid", fgColor="DBEAFE")
        summary_font = Font(bold=True, size=10, color="1E40AF")
        thin = Side(style="thin", color="D1D5DB")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)

        items = sorted(collection.template.items, key=lambda x: x.order)
        dept_map = {d["id"]: d["name"] for d in INITIAL_DEPARTMENTS}
        active_criminals = [c for c in collection.criminalists if c.is_active]

        # ── Ширины колонок и высоты строк ─────────────────────────
        # Название пункта должно помещаться в 2 строки без обрезки
        n_items = len(items)
        max_name_len = max((len(item.name) for item in items), default=10)
        item_width = min(30, max(12, int(-(-max_name_len // 2) * 1.15) + 2))
        name_w = 28 if n_items <= 4 else max(16, 28 - (n_items - 4) * 2)
        header_lines = max(1, -(-max_name_len // max(1, item_width - 2)))
        header_height = max(30, header_lines * 15 + 8)

        # ── Заголовок ─────────────────────────────────────────────
        row = 1
        total_cols = 2 + len(items) + 1  # № + ФИО + пункты + итог
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=total_cols)
        title_cell = ws.cell(row=row, column=1,
                             value=f"СВОДКА: {collection.template.name}")
        title_cell.fill = header_fill
        title_cell.font = Font(bold=True, color="FFFFFF", size=12)
        title_cell.alignment = center
        title_cell.border = border
        ws.row_dimensions[row].height = 24
        row += 1

        # Дата + статистика
        now_str = datetime.now().strftime("%d.%m.%Y %H:%M")
        done_count = sum(1 for c in active_criminals
                         if get_criminalist_fill(collection, c)["percent"] >= 100)
        partial_count = sum(1 for c in active_criminals
                            if 0 < get_criminalist_fill(collection, c)["percent"] < 100)
        not_done = len(active_criminals) - done_count - partial_count
        inactive_count = sum(1 for c in collection.criminalists if not c.is_active)

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=total_cols)
        info_text = (f"Дата: {now_str}  |  "
                     f"Активных: {len(active_criminals)}  |  "
                     f"Сдали: {done_count}  |  "
                     f"Частично: {partial_count}  |  "
                     f"Не сдали: {not_done}")
        if inactive_count:
            info_text += f"  |  Откл: {inactive_count}"
        info_cell = ws.cell(row=row, column=1, value=info_text)
        info_cell.fill = subheader_fill
        info_cell.font = Font(bold=True, color="FFFFFF", size=9)
        info_cell.alignment = center
        info_cell.border = border
        ws.row_dimensions[row].height = 18
        row += 1

        # ── Шапка таблицы ─────────────────────────────────────────
        headers = ["№", "ФИО криминалиста"]
        for item in items:
            # Полное название — переносится в 2 строки, не обрезается
            headers.append(item.name)
        headers.append("%")

        for col_i, h in enumerate(headers, 1):
            c = ws.cell(row=row, column=col_i, value=h)
            c.fill = header_fill
            c.font = Font(bold=True, color="FFFFFF", size=9)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = border
        ws.row_dimensions[row].height = header_height
        row += 1

        # ── Данные по криминалистам ───────────────────────────────
        # Сортируем: сдавшие → частично → не сдавшие
        def sort_key(c):
            fill = get_criminalist_fill(collection, c)
            pct = fill["percent"]
            if pct >= 100:
                return (0, -pct, c.full_name)
            elif pct > 0:
                return (1, -pct, c.full_name)
            else:
                return (2, 0, c.full_name)

        sorted_criminals = sorted(active_criminals, key=sort_key)

        for idx, crim in enumerate(sorted_criminals, 1):
            fill = get_criminalist_fill(collection, crim)
            pct = fill["percent"]

            # Определяем стиль строки
            if pct >= 100:
                row_fill, row_font = green_fill, green_font
            elif pct > 0:
                row_fill, row_font = yellow_fill, yellow_font
            else:
                row_fill, row_font = red_fill, red_font

            # Номер
            c1 = ws.cell(row=row, column=1, value=idx)
            c1.alignment = center
            c1.border = border
            c1.fill = row_fill
            c1.font = data_font

            # ФИО
            name = crim.full_name
            if crim.note:
                name += f" ({crim.note})"
            c2 = ws.cell(row=row, column=2, value=name)
            c2.alignment = left_wrap
            c2.border = border
            c2.fill = row_fill
            c2.font = name_font

            # Пункты
            for col_offset, item in enumerate(items, 3):
                item_filled = is_item_filled(collection, crim, item)

                if item.item_type == ReportItemType.NUMERICAL:
                    # Для числовых — показываем значение
                    rd = None
                    for sub in collection.submissions:
                        if (sub.criminalist_id == crim.id
                                and sub.template_item_id == item.id):
                            rd = sub
                            break

                    use_dept_mode = (collection.template.use_departments_mode
                                     and bool(crim.zone.department_ids))

                    if use_dept_mode:
                        departments = crim.zone.department_ids
                        if rd is None:
                            val_text = "—"
                        else:
                            total = sum(
                                v for v in (rd.department_values.get(did) for did in departments)
                                if v is not None
                            )
                            done_depts = sum(
                                1 for did in departments
                                if rd.department_values.get(did) is not None
                            )
                            val_text = f"{total} ({done_depts}/{len(departments)})"
                    else:
                        if rd and rd.value is not None:
                            unit = f" {item.unit}" if item.unit else ""
                            val_text = f"{rd.value}{unit}"
                        else:
                            val_text = "—"
                else:
                    val_text = "✓" if item_filled else "—"

                cell = ws.cell(row=row, column=col_offset, value=val_text)
                cell.alignment = center
                cell.border = border

                if item_filled:
                    cell.fill = green_fill
                    cell.font = Font(size=10, bold=True, color="15803D")
                else:
                    cell.fill = PatternFill("solid", fgColor="FFFFFF")
                    cell.font = Font(size=10, color="94A3B8")

            # Итоговый процент
            pct_col = total_cols
            pct_cell = ws.cell(row=row, column=pct_col, value=f"{pct}%")
            pct_cell.alignment = center
            pct_cell.border = border
            pct_cell.fill = row_fill
            pct_cell.font = Font(size=10, bold=True,
                                 color="15803D" if pct >= 100
                                 else "92400E" if pct > 0 else "991B1B")

            # Высота строки — под перенос длинного ФИО на 2 строки
            name_lines = max(1, -(-len(name) // max(1, name_w - 2)))
            ws.row_dimensions[row].height = max(20, name_lines * 15 + 6)
            row += 1

        # ── Сводка по пунктам ─────────────────────────────────────
        row += 1
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=total_cols)
        sum_cell = ws.cell(row=row, column=1, value="ИТОГИ ПО ПУНКТАМ")
        sum_cell.fill = header_fill
        sum_cell.font = header_font
        sum_cell.alignment = center
        sum_cell.border = border
        ws.row_dimensions[row].height = 20
        row += 1

        # Строка "Заполнили"
        c1 = ws.cell(row=row, column=1, value="")
        c1.border = border
        c2 = ws.cell(row=row, column=2, value="Заполнили:")
        c2.fill = summary_fill
        c2.font = summary_font
        c2.alignment = left_wrap
        c2.border = border

        total_active = len(active_criminals)
        for col_offset, item in enumerate(items, 3):
            filled_count = sum(1 for c in active_criminals
                               if is_item_filled(collection, c, item))

            if item.item_type == ReportItemType.NUMERICAL:
                # Для числовых — сумма
                total_val = 0
                for c in active_criminals:
                    for sub in collection.submissions:
                        if sub.criminalist_id == c.id and sub.template_item_id == item.id:
                            if sub.value is not None:
                                total_val += sub.value
                            elif sub.department_values:
                                total_val += sum(
                                    v for v in sub.department_values.values()
                                    if v is not None
                                )
                cell_val = f"{total_val} ({filled_count}/{total_active})"
            else:
                cell_val = f"{filled_count}/{total_active}"

            cell = ws.cell(row=row, column=col_offset, value=cell_val)
            cell.fill = summary_fill
            cell.font = summary_font
            cell.alignment = center
            cell.border = border

        pct_col = total_cols
        overall = round(done_count / max(total_active, 1) * 100)
        overall_cell = ws.cell(row=row, column=pct_col, value=f"{overall}%")
        overall_cell.fill = summary_fill
        overall_cell.font = Font(size=10, bold=True, color="1E40AF")
        overall_cell.alignment = center
        overall_cell.border = border
        ws.row_dimensions[row].height = 20

        # ── Ширина столбцов (вычислена выше, до шапки) ────────────
        ws.column_dimensions["A"].width = 4
        ws.column_dimensions["B"].width = name_w
        for col_offset in range(3, total_cols):
            col_letter = get_column_letter(col_offset)
            ws.column_dimensions[col_letter].width = item_width
        # Итог
        ws.column_dimensions[get_column_letter(total_cols)].width = 7

        # ── Настройки печати (1 страница, книжная) ────────────────
        ws.page_setup.orientation = "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = 0.4
        ws.page_margins.right = 0.4
        ws.page_margins.top = 0.5
        ws.page_margins.bottom = 0.5
        ws.page_margins.header = 0.2
        ws.page_margins.footer = 0.2

        # Закрепить шапку
        ws.freeze_panes = "A4"

        # Автофильтр
        ws.auto_filter.ref = f"A3:{get_column_letter(total_cols)}{3 + len(sorted_criminals)}"

        wb.save(filepath)
        print(f"[ZONAL_EXCEL] [OK] Файл сохранён: {filepath}")
