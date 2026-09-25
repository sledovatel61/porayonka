"""Генератор docx приказа из разметки черновика (@C, @R, @T, @TABLE, @SIGN, @PAGEBREAK, @N)."""
import sys, re
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn

TABLE_ROWS = [
    ("№ п/п", "Показатель", "Сведения"),
    ("1", "Количество уголовных дел, находящихся в производстве", ""),
    ("2", "Сроки предварительного следствия и даты их истечения", ""),
    ("3", "Принятые меры процессуального принуждения и обеспечительного характера", ""),
    ("4", "Количество уголовных дел, направленных прокурору с обвинительным заключением", ""),
    ("5", "Количество уголовных дел, направленных прокурором в суд", ""),
    ("6", "Количество прекращенных уголовных дел с указанием оснований принятия решений", ""),
    ("7", "Количество приостановленных уголовных дел с указанием оснований принятия решений", ""),
    ("8", "Количество уголовных дел, переданных по подследственности", ""),
    ("9", "Количество соединенных уголовных дел", ""),
    ("10", "Количество уголовных дел, возвращенных прокурором или судом для производства дополнительного следствия", ""),
    ("11", "Сведения о наиболее значимых уголовных делах, сроках исполнения поручений и имеющихся проблемных вопросах", ""),
]

def add_runs(p, text, size=14):
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        r = p.add_run(part[2:-2] if bold else part)
        r.bold = bold
        r.font.name = "Times New Roman"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        r.font.size = Pt(size)

def fmt(p, align, indent=True, before=0, after=0):
    pf = p.paragraph_format
    pf.alignment = align
    pf.first_line_indent = Cm(1.25) if indent else Cm(0)
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.0

def main(src, dst):
    d = docx.Document()
    s = d.sections[0]
    s.page_height, s.page_width = Cm(29.7), Cm(21.0)
    s.top_margin, s.bottom_margin = Cm(2), Cm(2)
    s.left_margin, s.right_margin = Cm(3), Cm(1.5)
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for line in open(src, encoding="utf-8").read().splitlines():
        if not line.strip():
            continue
        if line.startswith("@C "):
            p = d.add_paragraph(); add_runs(p, line[3:]); fmt(p, WD_ALIGN_PARAGRAPH.CENTER, False, 6, 6)
        elif line.startswith("@R "):
            p = d.add_paragraph(); add_runs(p, line[3:]); fmt(p, WD_ALIGN_PARAGRAPH.RIGHT, False)
        elif line.startswith("@T "):
            p = d.add_paragraph(); add_runs(p, "**" + line[3:] + "**"); fmt(p, WD_ALIGN_PARAGRAPH.CENTER, False, 12, 12)
        elif line.startswith("@N "):
            p = d.add_paragraph(); add_runs(p, line[3:], 12); fmt(p, WD_ALIGN_PARAGRAPH.LEFT, False, 6)
        elif line.strip() == "@PAGEBREAK":
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif line.strip() == "@SIGN":
            for t in ["", "", "Руководитель", "следственного управления"]:
                p = d.add_paragraph(); add_runs(p, t); fmt(p, WD_ALIGN_PARAGRAPH.LEFT, False)
        elif line.strip() == "@TABLE":
            t = d.add_table(rows=len(TABLE_ROWS), cols=3); t.style = "Table Grid"
            widths = [Cm(1.5), Cm(11.5), Cm(3.5)]
            for i, row in enumerate(TABLE_ROWS):
                for j, val in enumerate(row):
                    c = t.cell(i, j); c.width = widths[j]
                    p = c.paragraphs[0]; add_runs(p, ("**%s**" % val) if i == 0 and val else val, 12)
                    fmt(p, WD_ALIGN_PARAGRAPH.CENTER if j != 1 else WD_ALIGN_PARAGRAPH.LEFT, False)
        else:
            p = d.add_paragraph(); add_runs(p, line); fmt(p, WD_ALIGN_PARAGRAPH.JUSTIFY, True)
    d.save(dst)

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
